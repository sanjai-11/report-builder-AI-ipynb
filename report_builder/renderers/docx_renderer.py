from pathlib import Path

from docx import Document
from docx.shared import Inches

from ..models import OutputKind, ReportDocument


def render_docx(document: ReportDocument, destination: str | Path) -> Path:
    out_path = Path(destination)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    if document.config.logo_path and Path(document.config.logo_path).exists():
        doc.add_picture(str(document.config.logo_path), width=Inches(1.6))

    doc.add_heading(document.config.title, level=0)
    doc.add_paragraph(f"{document.config.company_name} | Audience: {document.config.audience.value}")

    if document.executive_summary:
        doc.add_heading("Executive Summary", level=1)
        doc.add_paragraph(document.executive_summary)

    for output in document.outputs:
        doc.add_heading(output.title, level=2)
        if output.kind is OutputKind.TABLE:
            rows = [line.split("|") for line in output.content.splitlines() if line.strip()]
            if rows:
                table = doc.add_table(rows=1, cols=len(rows[0]))
                table.style = "Light Grid Accent 1"
                header = table.rows[0].cells
                for i, value in enumerate(rows[0]):
                    header[i].text = value.strip()
                for row_values in rows[1:]:
                    row_cells = table.add_row().cells
                    for i, value in enumerate(row_values):
                        row_cells[i].text = value.strip()
            else:
                doc.add_paragraph(output.content)
        else:
            doc.add_paragraph(output.content)

    doc.save(str(out_path))
    return out_path
